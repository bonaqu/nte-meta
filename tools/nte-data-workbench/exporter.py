"""Lossless, categorized exports of the local NTE reference. No network access."""
from __future__ import annotations

import argparse
import csv
import gzip
import html
import json
import os
import re
import shutil
from collections import defaultdict
from pathlib import Path

KINDS = {
    'character': 'Персонажи', 'biography': 'Биографии', 'skill': 'Навыки',
    'awakening': 'Пробуждения', 'arc': 'Дуги', 'cartridge': 'Картриджи',
    'affinity': 'Симпатия', 'wardrobe': 'Гардероб', 'item': 'Предметы',
    'stat': 'Характеристики', 'image': 'Изображения', 'audio': 'Озвучка',
    'localization': 'Локализация',
}
COLUMNS = ['id', 'kind', 'ownerId', 'ownerName', 'namespace', 'gameId', 'ru', 'en', 'ja', 'zh', 'assets', 'details', 'jsonFile']

def export_kind(row):
    namespace=row['namespace'].lower()
    if row['id'].split('|')[-1].startswith('row:'):
        if 'fork' in namespace:return 'arc'
        if any(value in namespace for value in ['skill','gameplayability','characterability']):return 'skill'
        if 'effectfigure' in namespace or 'upgrade' in namespace:return 'stat'
    return row['kind']


def safe(value: str) -> str:
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', value).strip(' .')[:100] or 'unknown'


def records(repository: Path):
    with gzip.open(repository / 'scripts/content/game-data/chunks.jsonl.gz', 'rt', encoding='utf-8') as stream:
        for line in stream:
            yield from json.loads(line)['records']


def export(repository: Path, destination: Path, include_media: bool = True, progress=print):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    repository, destination = repository.resolve(), destination.resolve()
    if destination.exists() and any(destination.iterdir()):
        raise ValueError('Папка экспорта должна быть пустой: существующие данные не перезаписываются.')
    destination.mkdir(parents=True, exist_ok=True)
    owners = {};variants={}
    for row in records(repository):
        if row['namespace'] == 'ST_Player' and row['gameId'].startswith('Role_Name_'):
            owners[row['ownerId']] = row['values'].get('ru') or row['values'].get('en') or row['ownerId']
        elif row['namespace']=='NTE_ReferenceCharacter':owners.setdefault(row['ownerId'],row['values'].get('ru') or row['values'].get('en') or row['ownerId'])
        elif row['namespace']=='DT_DefaultAppearanceData':
            asset=row['details'].get('PlayerAppearanceAsset',{}).get('AssetPathName','').lower()
            if '_female/' in asset:variants[row['ownerId']]='женский вариант'
            elif '_male/' in asset:variants[row['ownerId']]='мужской вариант'
    for owner,variant in variants.items():owners[owner]=owners.get(owner,owner)+' ('+variant+')'
    workbook = Workbook(write_only=True)
    index = workbook.create_sheet('Содержание')
    index.append(['Раздел', 'Что хранится'])
    index.append(['JSON', 'Полные исходные значения, поля и связи без ограничения длины Excel'])
    index.append(['CSV', 'UTF-8 BOM; ключи и все четыре языка. Длинные поля находятся в JSON'])
    index.append(['Персонажи', 'Отдельная папка по игровому ID; принадлежность только из исходных связей'])
    index.append(['Общие данные', 'Записи без подтверждённого владельца остаются в общих разделах'])
    sheets, csv_files, json_files, owner_files = {}, {}, {}, {}
    owner_media=defaultdict(dict)
    counts = defaultdict(int)
    try:
        for number, row in enumerate(records(repository), 1):
            kind = export_kind(row)
            if kind not in sheets:
                title = KINDS.get(kind, kind)
                folder = destination / 'Категории' / safe(title)
                folder.mkdir(parents=True)
                sheets[kind] = workbook.create_sheet(title[:31])
                sheets[kind].append(COLUMNS)
                sheets[kind].freeze_panes = 'G2'
                csv_files[kind] = (folder / 'таблица.csv').open('w', encoding='utf-8-sig', newline='')
                csv.writer(csv_files[kind], delimiter=';').writerow(COLUMNS)
                json_files[kind] = (folder / 'данные.jsonl').open('w', encoding='utf-8')
                index.append([title, 'Категории/' + title + '/данные.jsonl'])
            counts[kind] += 1
            folder = destination / 'Категории' / safe(KINDS.get(kind, kind))
            line = json.dumps(row, ensure_ascii=False, separators=(',', ':'))
            json_files[kind].write(line + '\n')
            relative = str((folder / 'данные.jsonl').relative_to(destination)) + f' (строка {counts[kind]})'
            values = row['values']
            flat = [row['id'], row['kind'], row['ownerId'], owners.get(row['ownerId'], ''), row['namespace'], row['gameId'],
                    *[values.get(locale, '') for locale in ['ru', 'en', 'ja', 'zh']],
                    json.dumps(row['assets'], ensure_ascii=False), json.dumps(row['details'], ensure_ascii=False), relative]
            flat = [str(v) for v in flat]
            # Never truncate the only copy: the full value remains in JSONL.
            short = [v if len(v) <= 32000 else 'Полное значение: ' + relative for v in flat]
            cells = []
            for value in short:
                cell = WriteOnlyCell(sheets[kind], value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value))
                cell.data_type = 's'  # Game text must never become an Excel formula.
                cells.append(cell)
            sheets[kind].append(cells)
            csv.writer(csv_files[kind], delimiter=';').writerow(["'"+v if v.startswith(('=', '+', '-', '@')) else v for v in short])
            if row['ownerId']:
                owner = row['ownerId']
                if owner not in owner_files:
                    owner_dir = destination / 'Персонажи' / safe(owner + ' — ' + owners.get(owner, 'Без имени'))
                    owner_dir.mkdir(parents=True)
                    owner_files[owner] = {}
                owner_streams = owner_files[owner]
                if kind not in owner_streams:
                    owner_streams[kind] = (destination / 'Персонажи' / safe(owner + ' — ' + owners.get(owner, 'Без имени')) / (safe(KINDS.get(kind, kind)) + '.jsonl')).open('w', encoding='utf-8')
                owner_streams[kind].write(line + '\n')
                for asset in row.get('assets',[]):
                    url=asset.get('url','')
                    if url.startswith(('/nte-meta/game-assets/','/nte-meta/game-audio/')):
                        media_category='Озвучка' if kind=='audio' else 'Пробуждения' if 'juexingicon' in asset.get('path','').lower() else 'Навыки' if '/skill/' in asset.get('path','').lower() else KINDS.get(kind,kind)
                        label=row['values'].get('ru') or row['gameId']
                        owner_media[owner][url]=(safe(media_category),safe(label),url.removeprefix('/nte-meta/'))
            if number % 10000 == 0:
                progress(f'Экспортировано {number:,} записей…')
        for kind,sheet in sheets.items():sheet.auto_filter.ref=f'A1:M{counts[kind]+1}'
        workbook.save(destination / 'NTE — справочник.xlsx')
    finally:
        for file in [*csv_files.values(), *json_files.values(), *[s for group in owner_files.values() for s in group.values()]]:
            file.close()
    manifest = json.loads((repository / 'scripts/content/game-data/manifest.json').read_text(encoding='utf-8'))
    (destination / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    if include_media:
        for name in ['game-assets', 'game-audio']:
            source = repository / 'public' / name
            if source.exists():
                progress('Копирование ' + name)
                shutil.copytree(source, destination / 'Медиа' / name)
        for owner,media in owner_media.items():
            owner_dir=destination/'Персонажи'/safe(owner+' — '+owners.get(owner,'Без имени'))
            for category,label,relative in media.values():
                original=destination/'Медиа'/relative
                if not original.exists():continue
                target=owner_dir/'Медиа'/category/(label+'-'+original.name)
                target.parent.mkdir(parents=True,exist_ok=True)
                try:os.link(original,target)
                except OSError:shutil.copy2(original,target)
    for owner,streams in owner_files.items():
        owner_dir=destination/'Персонажи'/safe(owner+' — '+owners.get(owner,'Без имени'))
        links=''.join('<li><a href="'+html.escape(safe(KINDS.get(kind,kind))+'.jsonl',quote=True)+'">'+html.escape(KINDS.get(kind,kind))+'</a></li>' for kind in streams)
        (owner_dir/'Содержание.html').write_text('<!doctype html><html lang="ru"><meta charset="utf-8"><title>'+html.escape(owners.get(owner,owner))+'</title><h1>'+html.escape(owners.get(owner,owner))+' · '+owner+'</h1><ul>'+links+'</ul><p>Иконки и реплики находятся в подпапке «Медиа». Полные поля, переводы и связи — в JSONL; поиск и фильтрация — в общем Excel.</p></html>',encoding='utf-8')
    rows_html = ''.join(f'<tr><td>{html.escape(KINDS.get(k,k))}</td><td>{v:,}</td></tr>' for k,v in sorted(counts.items()))
    empty = ''.join('<li>' + html.escape(p) + '</li>' for p in manifest.get('emptyTables', []))
    document = f'''<!doctype html><html lang="ru"><meta charset="utf-8"><title>Экспорт NTE Meta</title>
<style>body{{font:16px system-ui;background:#202126;color:#eee;max-width:1000px;margin:30px auto;padding:20px}}a{{color:#68e0eb}}td,th{{padding:8px 20px;border-bottom:1px solid #555}}code{{word-break:break-all}}</style>
<h1>Данные установленного клиента NTE</h1><p>Записей: {sum(counts.values()):,}. Excel и CSV — для просмотра, JSONL — полные значения.</p>
<p><a href="NTE — справочник.xlsx">Открыть Excel</a></p><table><tr><th>Раздел</th><th>Записей</th></tr>{rows_html}</table>
<h2>Как читать записи</h2><p>ownerId — игровой ID персонажа. namespace — исходная таблица; gameId — ключ строки. RU/EN/JA/ZH — официальные значения. details.references содержит точные ссылки: таблица, строка и поле. assets.path — путь в архиве игры, assets.url — адрес на сайте. Медиа лежат в папке «Медиа» с теми же именами файлов.</p>
<p>Категория не всегда означает принадлежность персонажу: общие материалы, шаблоны и интерфейсная локализация остаются общими. Старые и действующие ID сохраняются отдельно. Таблица DT_DefaultAppearanceData помогает определить действующий ID.</p>
<h2>Что пока не прочитано</h2><p>Пустые таблицы перечислены ниже. Параметры вида {{0}} сохранены как исходные, если связь с числом не установлена. Отсутствие связей явно остаётся отсутствием данных.</p><ul>{empty}</ul></html>'''
    (destination / 'Открыть отчёт.html').write_text(document, encoding='utf-8')
    progress(f'Готово: {sum(counts.values()):,} записей; {destination}')
    return dict(counts)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--repository', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--no-media', action='store_true')
    args = parser.parse_args()
    export(args.repository, args.output, not args.no_media)
