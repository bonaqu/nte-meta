"""Player-facing offline handbook, PNG images and named MP3 voice records."""
from __future__ import annotations
from concurrent.futures import ThreadPoolExecutor
import hashlib
import html
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
from urllib.parse import quote

LANGUAGES = {'en': 'Английский', 'ja': 'Японский', 'ko': 'Корейский', 'zh': 'Китайский'}
TYPES = {
    'CONSUMABLE': 'Расходуемые предметы', 'FURNITURE': 'Дом и мебель', 'AUTO_USE': 'Автоматически используемые',
    'QUEST': 'Задания', 'IMPORTANT': 'Особые предметы', 'FASHION': 'Гардероб', 'LOCKER_KEY': 'Ключи от шкафчиков',
    'VEHICLE_SKIN': 'Транспорт — облики', 'SHOWCASE': 'Коллекции', 'COOK_MATERIAL': 'Кулинария — ингредиенты',
    'CHARACTER_FRAGMENT': 'Персонажи — фрагменты', 'CHARACTER_CARD': 'Персонажи — получение',
    'FIGURES': 'Коллекции — фигурки', 'FLOWER': 'Сад — цветы', 'Lottery': 'Розыгрыши', 'MUSIC': 'Музыка',
    'GLIDE': 'Планеры', 'COOK_BOOK': 'Кулинария — рецепты', 'LIKEABILITY_GIFT': 'Симпатия — подарки',
    'SHOP_POINT': 'Валюты и жетоны', 'DEFENDATHENA_CONSUMABLE': 'Предметы режима защиты',
    'DECORATION': 'Дом — украшения', 'BIG_FLOWER': 'Сад — крупные цветы', 'SEED': 'Сад — семена',
    'FRUIT': 'Сад — плоды', 'FISH_KEEP': 'Рыбалка', 'EXCHANGE_TICKET': 'Билеты обмена',
    'VEHICLE_ASSISTANT': 'Транспорт — помощники', 'DISCOUNT_CARD': 'Скидочные карты',
    'MOVABLEPLAYER': 'Музыкальные устройства', 'DIVINATION_CRYSTAL': 'Предсказания',
    'LANDSPACE_CIRCLE': 'Дом — ландшафт', 'LANDSPACE_CIRCLE_BIG': 'Дом — ландшафт',
    'LANDSPACE': 'Дом — ландшафт', 'LANDSPACE_BIG': 'Дом — ландшафт',
}
CSS = '''body{font:16px/1.6 system-ui,sans-serif;background:#202126;color:#eee;max-width:1100px;margin:auto;padding:24px}a{color:#70e4ed}h1,h2,h3{line-height:1.25}h2{margin-top:36px;border-bottom:2px solid #ed3d8d;padding-bottom:8px}nav{display:flex;flex-wrap:wrap;gap:12px}article{border-top:1px solid #52535c;margin:24px 0;padding-top:16px}table{border-collapse:collapse;width:100%;margin:12px 0}td,th{border:1px solid #565760;text-align:left;padding:9px;vertical-align:top}th{background:#303139}img{max-width:100%;max-height:560px;object-fit:contain}img.icon{width:72px;height:72px;float:left;margin:0 16px 8px 0}p{white-space:pre-line;overflow-wrap:anywhere}.muted,small{color:#c4c4ce}audio{display:block;max-width:100%;margin:12px 0}details{padding:10px;border:1px solid #52535c;margin:10px 0}.scroll{overflow:auto}input{font:inherit;padding:10px;background:#303139;color:#fff;border:1px solid #777;width:min(95%,550px)}.empty{color:#ccc}@media(max-width:600px){body{padding:14px}td,th{padding:6px;font-size:14px}}'''


def safe(value):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', str(value)).strip(' .')[:76].rstrip(' .') or 'Без названия'


def esc(value):
    return html.escape(str('' if value is None else value), quote=True)


def link(target, label):
    return '<a href="' + quote(str(target).replace('\\', '/'), safe='/#') + '">' + esc(label) + '</a>'


def page(path, title, body, back=''):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text('<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>' + esc(title) + '</title><style>' + CSS + '</style>' + (link(back, '← Содержание') if back else '') + '<h1>' + esc(title) + '</h1>' + body + '</html>', encoding='utf-8')


def table(headers, rows):
    return '<div class="scroll"><table><thead><tr>' + ''.join('<th>' + esc(h) + '</th>' for h in headers) + '</tr></thead><tbody>' + ''.join('<tr>' + ''.join('<td>' + esc(v) + '</td>' for v in row) + '</tr>' for row in rows) + '</tbody></table></div>'


def skill_cost_sections(option, profile, progression):
    """Named per-skill totals and destination levels, without losing numeric costs."""
    labels = {'melee': 'Базовая атака', 'skill': 'Навык', 'ultraskill': 'Сверхспособность',
              'qte': 'Навык поддержки', 'passive': 'Пассивный навык', 'city': 'Повседневный навык'}
    groups = {'combat': 'Боевые навыки', 'support': 'Пассивные навыки', 'city': 'Повседневные навыки'}
    normalize = lambda value: re.sub(r'[^\w]', '', str(value).lower().removesuffix('_name'))
    result = []
    occurrences = {}
    for skill in option.get('skills', []):
        kind = skill['type']; label = labels.get(kind, 'Навык')
        index = occurrences.get(kind, 0); occurrences[kind] = index + 1
        abilities = profile.get('abilities', [])
        matches = [row for row in abilities if row['type'] == label]
        ability = next((row for row in abilities if normalize(row.get('id')) == normalize(skill['id'])), None)
        if ability is None: ability = matches[index] if index < len(matches) else {}
        totals = {}; stages = []
        for stage in skill['stages']:
            for item_id, amount in zip(stage['items_id'], stage['amount']):
                item = progression['items'].get(item_id, {})
                if not item.get('name'): raise ValueError('Missing skill material name: ' + item_id)
                totals[item_id] = totals.get(item_id, 0) + amount
                stages.append((stage.get('fromLevel', int(stage['level']) - 1), stage['level'], item['name'], amount))
        result.append({'group': groups[skill['kind']], 'type': label, 'name': ability.get('name', label),
                       'iconUrl': ability.get('iconUrl', ''), 'stages': stages,
                       'totals': [(progression['items'][item]['name'], count, progression['items'][item].get('iconUrl', '')) for item, count in totals.items()]})
    return result


def category(row):
    key = row['gameId'].lower()
    for prefix, name in [
        ('characterupmaterial', 'Развитие/Опыт персонажа'), ('weaponupmaterial', 'Развитие/Опыт дуги'),
        ('equipmentupmaterial', 'Развитие/Опыт картриджей и модулей'),
        ('characterbreakmaterial', 'Развитие/Возвышение персонажа'),
        ('weaponbreakmaterial', 'Развитие/Возвышение дуги'),
        ('ordinarymonmaterial', 'Развитие/Материалы противников'),
        ('worldboss', 'Развитие/Материалы мировых боссов'),
        ('weeklycloneboss', 'Развитие/Материалы еженедельных боссов'),
        ('skillupmaterial', 'Развитие/Улучшение навыков'),
        ('worldmaterial', 'Развитие/Материалы открытого мира'),
        ('common_material', 'Развитие/Общие материалы'),
        ('food', 'Кулинария/Еда и напитки'), ('furniture', 'Дом и мебель'),
        ('specialgift', 'Симпатия/Особые подарки'), ('affection', 'Симпатия/Предметы симпатии'),
        ('fashion', 'Внешний вид/Одежда'), ('businesscard', 'Профиль/Визитки'),
        ('bussinesscard', 'Профиль/Визитки'), ('frame', 'Профиль/Рамки'),
        ('title', 'Профиль/Титулы'), ('avatar', 'Профиль/Аватары'),
        ('wallpaper', 'Внешний вид/Обои'), ('phonecase', 'Внешний вид/Чехлы телефона'),
        ('gold', 'Валюты/Жук-монеты'), ('fons', 'Валюты/Фонсы'), ('annulith', 'Валюты/Аннулиты'),
        ('giftpackage', 'Наборы и контейнеры/Подарочные наборы'),
        ('equipmentbox', 'Наборы и контейнеры/Экипировка'), ('characterbox', 'Наборы и контейнеры/Персонажи'),
        ('weaponbox', 'Наборы и контейнеры/Дуги'), ('dlc', 'Наборы и контейнеры/Дополнительные наборы'),
        ('test', 'Служебные/Тестовые записи клиента'), ('universal', 'Расходуемые/Универсальные предметы'),
        ('combat', 'Бой/Расходуемые предметы'), ('volleyball', 'Миниигры/Волейбол'),
        ('scratch', 'Миниигры/Скретч-карты'), ('dice', 'Миниигры/Кубики'),
        ('stamina', 'Расходуемые/Энергия'), ('citystamina', 'Город/Энергия'),
    ]:
        if key.startswith(prefix): return name
    if row['namespace'] == 'NTE_ReferenceCurrency': return 'Валюты и жетоны'
    tags = row['details'].get('tag_ids', [])
    for tag, name in [('FoodMaterial', 'Кулинария/Ингредиенты'), ('Gift', 'Симпатия/Подарки'), ('furniture', 'Дом и мебель'), ('DevelopMaterial', 'Развитие/Другие материалы')]:
        if tag in tags: return name
    return TYPES.get(row['details'].get('type_id', '').removeprefix('ITEM_TYPE_'), 'Прочие предметы')


def image_category(row):
    path = ' '.join(asset.get('path', '') for asset in row.get('assets', [])).lower()
    for marker, title in [('appearance/fashion', 'Гардероб'), ('avatarimage', 'Портреты и аватары'), ('juexing', 'Пробуждения'), ('/skill/', 'Навыки'), ('/fork/', 'Дуги'), ('equipment', 'Картриджи и модули'), ('/item/', 'Предметы'), ('attribute', 'Характеристики')]:
        if marker in path: return title
    return 'Общий интерфейс и прочие изображения'


def compile_data(repository, progress):
    """Cache depends on exact snapshot AND parser sources, never on an old export."""
    paths = ['scripts/content/game-data/manifest.json', 'src/worker.js', 'scripts/game-data/export-readable.mjs']
    key = hashlib.sha256(b''.join((repository / p).read_bytes() for p in paths)).hexdigest()
    cache = Path(os.environ.get('LOCALAPPDATA', str(Path.home()))) / 'NTEMetaWorkbench/readable-cache'
    cache.mkdir(parents=True, exist_ok=True)
    target = cache / (key + '.json')
    if target.exists():
        progress('Используем готовые досье для этого снимка игры.')
        return json.loads(target.read_text(encoding='utf-8'))
    node = shutil.which('node')
    if not node: raise RuntimeError('Не найден Node.js. Он нужен для точной расшифровки игровых навыков тем же кодом, что использует сайт.')
    temporary = cache / (key + '.partial.json')
    process = subprocess.Popen([node, 'scripts/game-data/export-readable.mjs', '--output', str(temporary)], cwd=repository, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding='utf-8', errors='replace', creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
    for line in process.stdout: progress(line.rstrip())
    if process.wait(): raise RuntimeError('Не удалось подготовить игровые досье. Технический экспорт сохранён; подробности — в журнале.')
    data = json.loads(temporary.read_text(encoding='utf-8'))
    temporary.replace(target)
    return data


class Media:
    def __init__(self, repository, root, enabled):
        self.repository, self.root, self.enabled = repository, root, enabled
        self.images = {}
        self.audio_jobs = []
        self.ffmpeg = shutil.which('ffmpeg')
        if enabled and not self.ffmpeg: raise RuntimeError('Для совместимых MP3 нужен установленный FFmpeg.')

    def image(self, url, folder, label, large=False):
        if not self.enabled or not url.startswith('/nte-meta/game-assets/'): return ''
        original = self.repository / 'public' / url.removeprefix('/nte-meta/')
        if not original.is_file(): raise FileNotFoundError('Не найдено изображение: ' + str(original))
        central = self.root / 'Медиа/Изображения' / (original.stem + '.png')
        if url not in self.images:
            from PIL import Image
            central.parent.mkdir(parents=True, exist_ok=True)
            if central.exists():
                with Image.open(central) as image: image.verify()
            else:
                with Image.open(original) as image: image.save(central, 'PNG')
            self.images[url] = central
        local = folder / 'Медиа/Изображения' / (safe(label) + '-' + original.stem[:6] + '.png')
        local.parent.mkdir(parents=True, exist_ok=True)
        if not local.exists():
            try: os.link(central, local)
            except OSError: shutil.copy2(central, local)
        relative = os.path.relpath(local, folder).replace('\\', '/')
        return '<a href="' + quote(relative) + '"><img ' + ('' if large else 'class="icon" ') + 'src="' + quote(relative) + '" alt="' + esc(label) + '" loading="lazy"></a>'

    def audio_path(self, row, folder):
        if not self.enabled or not row.get('audioUrl', '').startswith('/nte-meta/game-audio/'): return None
        stem = Path(row['audioUrl']).stem
        return folder / 'Медиа/Озвучка' / safe(row.get('language') or 'Язык не указан') / (safe(row['title'])[:46] + '-' + stem[:8] + '.mp3')

    def voice(self, row, folder, character):
        target = self.audio_path(row, folder)
        if target is None: return ''
        original = self.repository / 'public' / row['audioUrl'].removeprefix('/nte-meta/')
        if not original.is_file(): raise FileNotFoundError('Не найдена реплика: ' + str(original))
        target.parent.mkdir(parents=True, exist_ok=True)
        self.audio_jobs.append((original, target, row, character))
        relative = quote(os.path.relpath(target, folder).replace('\\', '/'))
        return '<audio controls preload="none" src="' + relative + '"></audio><p>' + link(os.path.relpath(target, folder), 'Открыть MP3 в проигрывателе') + '</p>'

    def convert_audio(self, progress):
        def convert(job):
            original, target, row, character = job
            if target.exists(): return
            result = subprocess.run([self.ffmpeg, '-nostdin', '-v', 'error', '-i', str(original), '-map', '0:a:0', '-c:a', 'libmp3lame', '-q:a', '4', '-metadata', 'title=' + row['title'], '-metadata', 'artist=' + character, '-metadata', 'album=NTE · ' + row.get('language', ''), '-metadata', 'comment=' + row.get('description', ''), str(target)], capture_output=True, timeout=60, creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
            if result.returncode: raise RuntimeError('Не удалось создать MP3: ' + str(target))
        total = len(self.audio_jobs)
        with ThreadPoolExecutor(max_workers=4) as pool:
            for number, _ in enumerate(pool.map(convert, self.audio_jobs), 1):
                if number % 100 == 0 or number == total: progress(f'Совместимые MP3: {number} из {total}')


def export_readable(repository, destination, owners, source_records, include_media, progress, data=None):
    from openpyxl import Workbook
    book = Workbook(write_only=True)
    try:
        return _export_readable(repository, destination, owners, source_records, include_media, progress, data, book)
    finally:
        for sheet in book.worksheets:
            if not sheet.closed: sheet.close()
        book.close()


def _export_readable(repository, destination, owners, source_records, include_media, progress, data, book):
    from openpyxl.cell import WriteOnlyCell
    root = destination / 'Для игрока'
    root.mkdir(parents=True, exist_ok=True)
    data = compile_data(repository, progress) if data is None else data
    media = Media(repository, root, include_media)
    dictionary = {}; by_key = {}; asset_names = {}
    item_rows = []; image_rows = []; audio_rows = []; voice_parts = []
    for row in source_records():
        ru = row['values'].get('ru')
        if ru:
            by_key[row['namespace'] + '::' + row['gameId']] = ru
            for value in row['values'].values():
                dictionary.setdefault(value, set()).add(ru)
            if row.get('kind') != 'localization':
                for asset in row.get('assets', []): asset_names.setdefault(asset.get('url'), ru)
        if row['namespace'] in ('NTE_ReferenceItem', 'NTE_ReferenceCurrency'): item_rows.append(row)
        if row.get('kind') == 'image': image_rows.append(row)
        if row.get('kind') == 'audio': audio_rows.append(row)
        if re.search(r'FileDataTable\d*$', row['namespace']):
            voice_parts.extend(row['details'].get('DailyContent', [])); voice_parts.extend(row['details'].get('BattleContent', []))
    def native_text(value):
        if not isinstance(value, dict): return value or ''
        namespace = value.get('Namespace', value.get('TableId', '').split('.')[-1])
        return by_key.get(namespace + '::' + value.get('Key', ''), value.get('CultureInvariantString', ''))
    voice_labels = {r.get('AudioEvent', {}).get('AssetPathName', '').split('.')[-1]: (native_text(r.get('Title')), native_text(r.get('Content'))) for r in voice_parts}
    def translated(value):
        matches = dictionary.get(value, set())
        return re.sub(r'<[^>]+>', '', next(iter(matches))).strip() if len(matches) == 1 else ''
    sheets = {}; sheet_sizes = {}
    def excel(sheet, headers, values):
        if sheet not in sheets:
            sheets[sheet] = book.create_sheet(sheet)
            sheets[sheet].append(headers)
            sheets[sheet].freeze_panes = 'A2'
            sheet_sizes[sheet] = [1, len(headers)]
        cells = []
        for column, value in enumerate(values):
            text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str('' if value is None else value))
            if len(text) > 32000: raise ValueError('Игровая строка слишком длинна для читаемого Excel; полная копия сохранена в технических данных.')
            cell = WriteOnlyCell(sheets[sheet], text); cell.data_type = 's'
            if headers[column] in ('Досье', 'MP3') and text:
                cell.hyperlink = quote(text.replace('\\', '/'), safe='/'); cell.style = 'Hyperlink'
            cells.append(cell)
        sheets[sheet].append(cells)
        sheet_sizes[sheet][0] += 1
    character_links = []; covered_voices = set()
    sections = [('Биография', 'biography'), ('Характеристики', 'baseStats'), ('Навыки', 'abilities'), ('Пробуждения и резонанс', 'awakenings'), ('Прокачка', 'materials'), ('Симпатия и подарки', 'friendship'), ('Гардероб', 'skins'), ('Озвучка', 'voiceLines'), ('Рекомендации игры', 'recommendations')]
    for character in data['characters']:
        owner, profile = character['ownerId'], character['profile']
        name = owners.get(owner, character['name'])
        folder = root / 'Персонажи' / safe(owner + ' — ' + name)
        folder.mkdir(parents=True, exist_ok=True)
        (folder / 'Досье.json').write_text(json.dumps({
            'ownerId': owner, 'name': name, 'values': character['values'],
            'profile': profile, 'recommendations': character.get('recommendations') or {},
        }, ensure_ascii=False, indent=2), encoding='utf-8')
        progress('Понятные разделы: ' + name)
        character_links.append(link(os.path.relpath(folder / 'Содержание.html', root), name))
        index = '<p>Выберите раздел. Картинки можно открыть целиком; реплики прослушиваются на странице или в обычном проигрывателе.</p><nav>'
        index += ''.join(link(f'{i:02} — {title}.html', title) for i, (title, _) in enumerate(sections, 1)) + link('Досье.json', 'Данные JSON') + '</nav>'
        index += media.image(character.get('fields', {}).get('imageUrl', ''), folder, 'Портрет — ' + name)
        passport = [(label, profile.get(key)) for key, label in [('birthday', 'День рождения'), ('faction', 'Организация'), ('arcType', 'Тип дуги')] if profile.get(key)]
        if passport: index += table(['О персонаже', 'Значение'], passport)
        index += table(['Язык', 'Имя в игре'], [('Русский', character['values'].get('ru', '')), ('Английский', character['values'].get('en', '')), ('Японский', character['values'].get('ja', '')), ('Китайский', character['values'].get('zh', ''))])
        excel('Персонажи', ['Персонаж', 'Английское имя', 'Японское имя', 'Китайское имя', 'Досье'], [name, character['values'].get('en'), character['values'].get('ja'), character['values'].get('zh'), os.path.relpath(folder / 'Содержание.html', root)])
        for number, (title, field) in enumerate(sections, 1):
            body = ''
            if field == 'biography':
                body = '<p>' + esc(profile.get('biography') or profile.get('biographyShort') or '') + '</p>'
                if passport: body += table(['О персонаже', 'Значение'], passport)
                if profile.get('voiceActors'): body += '<h2>Актёры озвучки</h2>' + table(['Язык', 'Актёр'], [(r['language'], r['name']) for r in profile['voiceActors']])
            elif field == 'baseStats':
                body = '<p>Начальные показатели персонажа. Бонусы экипировки и отряда сюда не включены.</p>' + table(['Показатель', 'Значение'], [(r['label'], r['value']) for r in profile.get(field, [])])
                for r in profile.get(field, []): excel('Характеристики', ['Персонаж', 'Показатель', 'Начальное значение'], [name, r['label'], r['value']])
            elif field == 'recommendations':
                recommendations = character.get('recommendations') or {}
                body = '<h2>Дуги</h2>'
                for i, arc in enumerate(recommendations.get('arcs', [])):
                    label = 'Основной вариант' if i == 0 else 'На замену'
                    body += '<article>' + media.image(arc.get('imageUrl', ''), folder, 'Рекомендуемая дуга — ' + arc['name']) + '<h3>' + esc(label + ' · ' + arc['name']) + '</h3><p>' + esc(arc.get('description')) + '</p>' + table(['Характеристика', 'Значение'], [(s['label'], s['value']) for s in arc.get('stats', [])]) + '</article>'
                    excel('Рекомендации дуг', ['Персонаж', 'Вариант', 'Дуга', 'Эффект'], [name, label, arc['name'], arc.get('description')])
                body += '<h2>Основные характеристики</h2>'
                for stat in recommendations.get('stats', []):
                    body += '<article>' + media.image(stat.get('imageUrl', ''), folder, 'Рекомендуемая характеристика — ' + stat['name']) + '<h3>' + esc(stat['name']) + '</h3></article>'
                    excel('Рекомендуемые статы', ['Персонаж', 'Характеристика'], [name, stat['name']])
                body += '<h2>Рекомендуемые уровни навыков</h2><p>Рекомендации клиента для разных этапов развития. Числа обозначают целевые уровни навыков, а не их приоритет.</p>'
                for stage in recommendations.get('stages', []):
                    body += '<h3>Этап ' + esc(stage['stage']) + ' · персонаж ур. ' + esc(stage.get('characterLevel')) + ' · дуга ур. ' + esc(stage.get('arcLevel')) + '</h3>' + table(['Тип навыка', 'Навык', 'Рекомендуемый уровень'], [(r['skillType'], r['name'], r['level']) for r in stage.get('skills', [])])
                    for skill in stage.get('skills', []):
                        excel('Рекомендуемые навыки', ['Персонаж', 'Этап', 'Уровень персонажа', 'Уровень дуги', 'Тип навыка', 'Навык', 'Уровень навыка'], [name, stage['stage'], stage.get('characterLevel'), stage.get('arcLevel'), skill['skillType'], skill['name'], skill['level']])
                if not recommendations:
                    body += '<p>В этом снимке рекомендации персонажа отсутствуют.</p>'
            elif field == 'abilities':
                groups = {}
                for r in profile.get(field, []): groups.setdefault(r['type'], []).append(r)
                for kind, rows in groups.items():
                    body += '<h2>' + esc(kind) + '</h2>'
                    for r in rows:
                        body += '<article>' + media.image(r.get('iconUrl', ''), folder, 'Навык — ' + r['name']) + '<h3>' + esc(r['name']) + '</h3><p>' + esc(r.get('description')) + '</p>'
                        excel('Навыки', ['Персонаж', 'Тип навыка', 'Название', 'Описание'], [name, kind, r['name'], r.get('description')])
                        for attribute in r.get('attributes', []):
                            levels = re.split(r'\s*·\s*(?=\d+:)', attribute['value'])
                            body += '<h4>' + esc(attribute['label']) + '</h4>' + table(['Уровень навыка', 'Значение'], [((v.split(':', 1)[0], v.split(':', 1)[1]) if re.match(r'^\d+:', v) else ('Все уровни', v)) for v in levels])
                            for v in levels: excel('Параметры навыков', ['Персонаж', 'Навык', 'Параметр', 'Уровень и значение'], [name, r['name'], attribute['label'], v])
                        body += '</article>'
            elif field == 'awakenings':
                for kind, label in [('awakening', 'Пробуждения'), ('resonance', 'Бонус за 3 / 6 пробуждений')]:
                    body += '<h2>' + label + '</h2>'
                    for r in profile.get(field, []):
                        if r.get('kind', 'awakening') != kind: continue
                        body += '<article>' + media.image(r.get('iconUrl', ''), folder, label + ' — ' + r['name']) + '<h3>' + esc(str(r['level']) + ' · ' + r['name']) + '</h3><p>' + esc(r.get('description')) + '</p></article>'
                        excel('Пробуждения', ['Персонаж', 'Раздел', 'Число пробуждений', 'Название', 'Эффект'], [name, label, r['level'], r['name'], r.get('description')])
            elif field == 'materials':
                option = next((c for c in data.get('progression', {}).get('costs', []) if c['kind'] == 'character' and c['id'] == owner), None)
                if option:
                    body += '<p>Возвышение повышает предельный уровень. Например, возвышение на 20-м уровне нужно для перехода с 20 на 21 и открывает уровни до 30.</p>'
                    for stage in option['stages']:
                        rows = [(data['progression']['items'].get(item, {}).get('name', item), amount) for item, amount in zip(stage['items_id'], stage['amount'])]
                        if stage.get('cost_gold'): rows.append(('Жук-монета', stage['cost_gold']))
                        body += '<h2>На уровне ' + esc(stage.get('fromLevel')) + ' → предел ' + esc(stage['level']) + '</h2>' + table(['Материал', 'Нужно'], rows)
                        for item, amount in rows: excel('Возвышение', ['Персонаж', 'На уровне', 'Новый предел', 'Материал', 'Количество'], [name, stage.get('fromLevel'), stage['level'], item, amount])
                body += '<h2>Всего на возвышения</h2>'
                for r in profile.get(field, []): body += '<article>' + media.image(r.get('iconUrl', ''), folder, r['name']) + '<h3>' + esc(r['name']) + '</h3><p>' + esc(r.get('amount', r.get('quantity'))) + '</p><p>' + esc(r.get('source')) + '</p></article>'
                for skill in skill_cost_sections(option or {}, profile, data.get('progression', {})):
                    body += '<article><h2>' + esc(skill['group'] + ' · ' + skill['type']) + '</h2>' + media.image(skill['iconUrl'], folder, 'Навык — ' + skill['name']) + '<h3>' + esc(skill['name']) + '</h3><h4>Всего для этого навыка</h4>'
                    for item, amount, icon in skill['totals']:
                        body += '<article>' + media.image(icon, folder, item) + '<h4>' + esc(item) + ' ×' + esc(amount) + '</h4></article>'
                        excel('Итоги навыков', ['Персонаж', 'Раздел', 'Тип навыка', 'Навык', 'Материал', 'Количество'], [name, skill['group'], skill['type'], skill['name'], item, amount])
                    body += '<details><summary>Материалы по уровням</summary>' + table(['С уровня', 'До уровня', 'Материал', 'Количество'], skill['stages']) + '</details></article>'
                    for start, end, item, amount in skill['stages']:
                        excel('Прокачка навыков', ['Персонаж', 'Раздел', 'Навык', 'С уровня', 'До уровня', 'Материал', 'Количество'], [name, skill['group'], skill['name'], start, end, item, amount])
                body += '<details><summary>Опыт для каждого следующего уровня</summary>' + table(['Текущий уровень', 'Опыт до следующего'], [(r['level'], r['exp']) for r in data.get('progression', {}).get('characterExp', []) if r['level'] < 80]) + '</details>'
            elif field == 'friendship':
                body += '<h2>Уровни симпатии и награды</h2>'
                for r in profile.get('friendship', []):
                    rewards = [(x['name'], x.get('quantity', '')) for x in r.get('rewards', [])]
                    body += '<article><h3>Уровень ' + esc(r['level']) + '</h3><p>' + esc(r.get('description')) + '</p>' + table(['Награда', 'Количество'], rewards) + '</article>'
                    for item, amount in rewards: excel('Симпатия', ['Персонаж', 'Уровень', 'Награда', 'Количество', 'Опыт'], [name, r['level'], item, amount, r.get('description')])
                body += '<h2>Любимые подарки</h2><p>Очки симпатии указаны для этого персонажа.</p>'
                for r in profile.get('gifts', []):
                    body += '<article>' + media.image(r.get('iconUrl', ''), folder, r['name']) + '<h3>' + esc(r['name']) + '</h3><p>' + esc(r.get('effect')) + '</p></article>'
                    excel('Подарки', ['Персонаж', 'Подарок', 'Очки симпатии'], [name, r['name'], r.get('effect')])
            elif field == 'skins':
                for r in profile.get(field, []):
                    body += '<article><h2>' + esc(r['name']) + '</h2>' + media.image(r.get('imageUrl', ''), folder, 'Гардероб — ' + r['name'], True) + '<p>' + esc(r.get('description')) + '</p></article>'
                    excel('Гардероб', ['Персонаж', 'Название', 'Описание'], [name, r['name'], r.get('description')])
            elif field == 'voiceLines':
                for language in LANGUAGES.values():
                    rows = [r for r in profile.get(field, []) if r.get('language') == language and r.get('audioUrl')]
                    body += '<h2>' + language + '</h2>'
                    if not rows: body += '<p>В этом снимке нет доступных звуковых файлов этого языка.</p>'
                    for r in rows:
                        covered_voices.add((owner, r.get('id')))
                        body += '<article><h3>' + esc(r['title']) + '</h3><p>' + esc(r.get('description')) + '</p>' + media.voice(r, folder, name) + '</article>'
                        audio_path = media.audio_path(r, folder)
                        excel('Реплики', ['Персонаж', 'Язык', 'Ситуация', 'Текст реплики', 'MP3'], [name, language, r['title'], r.get('description'), os.path.relpath(audio_path, root) if audio_path else ''])
            if not body.strip() or body == '<p></p>': body = '<p>В этом снимке нет подтверждённых данных раздела.</p>'
            page(folder / f'{number:02} — {title}.html', name + ' · ' + title, body, 'Содержание.html')
        page(folder / 'Содержание.html', name, index, '../../Начать здесь.html')
    equipment_links = []
    for kind, title in [('arcs', 'Дуги'), ('cartridges', 'Картриджи и консоли')]:
        body = ''
        folder = root / 'Экипировка' / title
        for r in data.get('equipment', []):
            if r['kind'] != kind: continue
            body += '<article><h2>' + esc(r['name']) + '</h2>' + media.image(r.get('imageUrl', ''), folder, r['name']) + '<p>' + esc(r.get('description')) + '</p>' + table(['Характеристика', 'Значение'], [(s['label'], s['value']) for s in r.get('stats', [])])
            excel(title[:31], ['Название', 'Описание', 'Характеристики'], [r['name'], r.get('description'), '\n'.join(s['label'] + ': ' + s['value'] for s in r.get('stats', []))])
            for index, stat in enumerate(r.get('details', {}).get('stats', [])):
                label = r['stats'][index]['label'].split(' (ур.')[0] if index < len(r.get('stats', [])) else stat.get('name', 'Показатель')
                values = [(level, str(value) + ('%' if stat.get('b_is_percent') else '')) for level, value in enumerate(stat.get('values', []), 1)]
                body += '<details><summary>' + esc(label) + ' · все уровни</summary>' + table(['Уровень дуги', 'Значение'], values) + '</details>'
                for level, value in values: excel('Дуги — уровни', ['Дуга', 'Характеристика', 'Уровень', 'Значение'], [r['name'], label, level, value])
            option = next((c for c in data.get('progression', {}).get('costs', []) if c['kind'] == 'arc' and c['id'] == r.get('gameId')), None)
            if option:
                body += '<details><summary>Материалы возвышения дуги</summary>'
                for stage in option['stages']:
                    amounts = [(data['progression']['items'].get(item, {}).get('name', item), amount) for item, amount in zip(stage['items_id'], stage['amount'])]
                    if stage.get('cost_gold'): amounts.append(('Жук-монета', stage['cost_gold']))
                    body += '<h3>На уровне ' + esc(stage.get('fromLevel')) + ' → предел ' + esc(stage['level']) + '</h3>' + table(['Материал', 'Нужно'], amounts)
                    for item, amount in amounts: excel('Возвышение дуг', ['Дуга', 'На уровне', 'Новый предел', 'Материал', 'Количество'], [r['name'], stage.get('fromLevel'), stage['level'], item, amount])
                body += '</details>'
            body += '</article>'
        page(folder / 'Каталог.html', title, body, '../../Начать здесь.html')
        equipment_links.append(link(os.path.relpath(folder / 'Каталог.html', root), title))
    groups = {}
    for r in item_rows: groups.setdefault(category(r), []).append(r)
    item_links = []
    for group, rows in sorted(groups.items()):
        folder = root / 'Предметы' / Path(*[safe(part) for part in group.split('/')])
        body = '<p>Поиск по этому разделу: Ctrl+F. Количество предметов: ' + str(len(rows)) + '.</p>'
        for r in sorted(rows, key=lambda row: row['values'].get('ru', row['gameId'])):
            name = r['values'].get('ru') or r['values'].get('en') or r['gameId']
            description = '\n\n'.join(dict.fromkeys(filter(None, [translated(r['details'].get('context', '')), translated(r['details'].get('description', ''))])))
            body += '<article><h2>' + esc(name) + '</h2>' + media.image(next((a['url'] for a in r.get('assets', []) if a['url'].startswith('/nte-meta/game-assets/')), ''), folder, name) + '<p>' + esc(description) + '</p><details><summary>Названия на других языках и игровой ID</summary>' + table(['RU', 'EN', 'JP', 'CN', 'ID'], [[r['values'].get(k, '') for k in ['ru', 'en', 'ja', 'zh']] + [r['gameId']]]) + '</details></article>'
            excel('Предметы', ['Раздел', 'Название', 'Описание', 'Английское имя', 'Японское имя', 'Китайское имя'], [group, name, description, r['values'].get('en'), r['values'].get('ja'), r['values'].get('zh')])
        page(folder / 'Каталог.html', group, body, os.path.relpath(root / 'Начать здесь.html', folder))
        item_links.append(link(os.path.relpath(folder / 'Каталог.html', root), group + ' · ' + str(len(rows))))
    # Keep all extracted voices, including old IDs and entries absent from a current profile.
    extra_voices = {}
    for r in audio_rows:
        if (r.get('ownerId'), r['details'].get('locale', '') + '-' + r['gameId']) not in covered_voices:
            extra_voices.setdefault(r.get('ownerId', ''), []).append(r)
    archive_links = []
    for owner, rows in extra_voices.items():
        name = owners.get(owner, 'Персонаж без подтверждённого имени · ID ' + owner)
        folder = root / 'Архив озвучки' / safe((owner + ' — ' + name)[:42])
        body = '<p>Эти файлы связаны с игровым ID, но не вошли в досье действующего персонажа. Это могут быть старые или дополнительные записи клиента. Неизвестная ситуация явно остаётся неподписанной.</p>'
        for locale, language in LANGUAGES.items():
            body += '<h2>' + language + '</h2>'
            for index, r in enumerate((r for r in rows if r['details'].get('locale') == locale), 1):
                title, description = voice_labels.get(r['gameId'], ('', ''))
                title = title or 'Реплика без игрового названия ' + str(index)
                voice = {'title': title, 'description': description, 'language': language, 'audioUrl': r['assets'][0]['url']}
                body += '<article><h3>' + esc(title) + '</h3><p>' + esc(description) + '</p>' + media.voice(voice, folder, name) + '<details><summary>Привязка в игре</summary><p>Игровой ID персонажа: ' + esc(owner) + '\nСобытие: ' + esc(r['gameId']) + '</p></details></article>'
                audio_path = media.audio_path(voice, folder)
                excel('Дополнительные реплики', ['Персонаж или ID', 'Язык', 'Ситуация', 'Текст реплики', 'Событие игры', 'MP3'], [name, language, title, description, r['gameId'], os.path.relpath(audio_path, root) if audio_path else ''])
        page(folder / 'Каталог.html', name + ' · дополнительные реплики', body, '../../Начать здесь.html')
        archive_links.append(link(os.path.relpath(folder / 'Каталог.html', root), name + ' · ' + str(len(rows))))
    gallery_links = []
    if include_media:
        galleries = {}
        for r in image_rows: galleries.setdefault(image_category(r), []).append(r)
        for group, rows in galleries.items():
            progress('PNG-медиатека: ' + group)
            folder = root / 'Медиатека' / group
            count = (len(rows) + 99) // 100
            def gallery_file(number): return 'Каталог.html' if number == 0 else f'Страница {number+1}.html'
            navigation = '<nav>' + ''.join(link(gallery_file(number), str(number+1)) for number in range(count)) + '</nav>'
            for number in range(count):
                body = '<p>Страница ' + str(number+1) + ' из ' + str(count) + '. По 100 изображений на странице.</p>' + navigation
                for index, r in enumerate(rows[number*100:(number+1)*100], number*100+1):
                    url = r['assets'][0]['url']
                    name = asset_names.get(url) or 'Изображение ' + str(index)
                    owner_name = owners.get(r.get('ownerId'), '')
                    body += '<article><h2>' + esc(name) + '</h2>' + ('<p>Персонаж: ' + esc(owner_name) + '</p>' if owner_name else '') + media.image(url, folder, name, True) + '<details><summary>Исходный путь в игре</summary><p>' + esc(r['assets'][0].get('path')) + '</p></details></article>'
                page(folder / gallery_file(number), group, body + navigation, '../../Начать здесь.html')
            gallery_links.append(link(os.path.relpath(folder / 'Каталог.html', root), group + ' · ' + str(len(rows))))
    # Write-only mode keeps the full technical dump and this concise workbook affordable.
    from openpyxl.utils import get_column_letter
    for title, sheet in sheets.items(): sheet.auto_filter.ref = f'A1:{get_column_letter(sheet_sizes[title][1])}{sheet_sizes[title][0]}'
    book.save(root / 'NTE — для игрока.xlsx')
    body = '<p>Это читаемый каталог установленного клиента. Начните с персонажа или типа предмета. Полные технические таблицы находятся в родительской папке.</p>'
    body += '<p>' + link('NTE — для игрока.xlsx', 'Открыть понятный Excel') + '</p><h2>Персонажи</h2><nav>' + ''.join(character_links) + '</nav><h2>Экипировка</h2><nav>' + ''.join(equipment_links) + '</nav><h2>Предметы по назначению</h2><nav>' + ''.join(item_links) + '</nav>'
    if gallery_links: body += '<h2>Полная медиатека изображений PNG</h2><nav>' + ''.join(gallery_links) + '</nav>'
    if archive_links: body += '<h2>Дополнительная и архивная озвучка</h2><nav>' + ''.join(archive_links) + '</nav>'
    body += '<h2>Изображения и озвучка</h2><p>Картинки сохранены в PNG. В досье откройте раздел «Озвучка»: реплики сгруппированы по языку и подписаны игровой ситуацией, например «Подарок получен» или «Солнечно». MP3 открываются в обычных проигрывателях. Внутренние WEM из рабочей выгрузки не предназначены для прослушивания.</p>' if include_media else '<p>Медиа выключены при экспорте: в этом каталоге только текст и таблицы.</p>'
    body += '<h2>Границы данных</h2><p>Отсутствующие файлы и нераскрытые связи клиента не выдумываются. Без установленного голосового пакета реплик этого языка не будет. Внешние дополнения отделены в исходных технических записях.</p>'
    page(root / 'Начать здесь.html', 'NTE · каталог для игрока', body)
    media.convert_audio(progress)
    progress('Читаемый каталог готов: ' + str(root / 'Начать здесь.html'))
    return {'characters': len(data['characters']), 'equipment': len(data.get('equipment', [])), 'items': len(item_rows), 'mp3': len(media.audio_jobs), 'png': len(media.images)}
