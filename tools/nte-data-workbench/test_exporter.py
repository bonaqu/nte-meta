import gzip
import json
from pathlib import Path
import tempfile
import unittest

from exporter import export


class ExportTests(unittest.TestCase):
    def test_lossless_export_formula_safety_and_character_media(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary);repo=root/'repository';out=root/'export'
            content=repo/'scripts/content/game-data';content.mkdir(parents=True)
            media=repo/'public/game-assets';media.mkdir(parents=True)
            (media/'test.webp').write_bytes(b'test-image')
            rows=[{'id':'character:00|loc:test','kind':'character','ownerId':'1025','namespace':'ST_Player','gameId':'Role_Name_025','values':{'ru':'Хатор','en':'Hathor','ja':'ハソール','zh':'哈索尔'},'assets':[],'details':{}},
                  {'id':'localization:00|row:test','kind':'localization','ownerId':'1025','namespace':'DT_SkillDamageData','gameId':'=danger','values':{'ru':'=SUM(1,2)'},'assets':[{'url':'/nte-meta/game-assets/test.webp','path':'HT/Content/Skill/test.uasset'}],'details':{'long':'я'*33000}}]
            with gzip.open(content/'chunks.jsonl.gz','wt',encoding='utf-8') as stream:stream.write(json.dumps({'records':rows},ensure_ascii=False)+'\n')
            (content/'manifest.json').write_text(json.dumps({'records':2,'emptyTables':[]}),encoding='utf-8')
            counts=export(repo,out,True,lambda _:None,readable=False)
            self.assertEqual(counts,{'character':1,'skill':1})
            workbook=load_workbook(out/'NTE — справочник.xlsx',read_only=True)
            cells=list(workbook['Навыки'].iter_rows())[1]
            self.assertEqual(cells[5].value,'=danger');self.assertEqual(cells[5].data_type,'s')
            self.assertIn('Полное значение',cells[11].value)
            workbook.close()
            recovered=json.loads((out/'Категории/Навыки/данные.jsonl').read_text(encoding='utf-8'))
            self.assertEqual(recovered,rows[1])
            self.assertIn("'=danger",(out/'Категории/Навыки/таблица.csv').read_text(encoding='utf-8-sig'))
            images=list((out/'Персонажи/1025 — Хатор/Медиа').rglob('*.webp'))
            self.assertEqual(len(images),1);self.assertEqual(images[0].read_bytes(),b'test-image')
            self.assertTrue((out/'Персонажи/1025 — Хатор/Содержание.html').exists())
            with self.assertRaises(ValueError):export(repo,out,False,lambda _:None)


if __name__=='__main__':unittest.main()
